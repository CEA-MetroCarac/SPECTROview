import os
import pytest
import numpy as np
import pandas as pd
from unittest.mock import MagicMock

from spectroview.model.m_settings import MSettings
from spectroview.viewmodel.vm_workspace_spectra import VMWorkspaceSpectra
from spectroview.viewmodel.utils import save_df_to_excel


@pytest.fixture
def test_df():
    return pd.DataFrame({
        'Filename': ['map1_10.5_zoneA', 'map1_20.3_zoneB', 'map1_30.1_zoneC'],
        'X': [1.0, 2.0, 3.0],
        'Y': [4.0, 5.0, 6.0],
        'P1_x0': [520.0, 521.0, 522.0],
        'P1_ampli': [10.0, 20.0, 30.0],
        'P2_ampli': [5.0, 10.0, 15.0],
    })


def test_save_df_to_excel(test_df, tmp_path):
    out_file = str(tmp_path / "test_results.xlsx")
    ok, msg = save_df_to_excel(out_file, test_df)
    assert ok is True
    assert os.path.exists(out_file)

    # Verify openpyxl can read it
    import openpyxl
    wb = openpyxl.load_workbook(out_file)
    ws = wb['Results']
    assert ws['A1'].value == 'Filename'
    assert ws['D1'].value == 'P1_x0'
    assert ws['B2'].value == 1.0
    wb.close()


def test_save_df_to_excel_empty(tmp_path):
    out_file = str(tmp_path / "empty.xlsx")
    ok, msg = save_df_to_excel(out_file, pd.DataFrame())
    assert ok is False


def test_add_column_from_filename(test_df, qapp):
    settings = MSettings()
    vm = VMWorkspaceSpectra(settings)
    vm.df_fit_results = test_df.copy()

    # Split part 1 ("10.5", "20.3", "30.1") which should convert to numeric floats
    vm.add_column_from_filename("Power", 1)
    assert "Power" in vm.df_fit_results.columns
    assert pd.api.types.is_numeric_dtype(vm.df_fit_results["Power"])
    assert list(vm.df_fit_results["Power"]) == [10.5, 20.3, 30.1]

    # Split part 2 ("zoneA", "zoneB", "zoneC") which should remain text
    vm.add_column_from_filename("Zone", 2)
    assert "Zone" in vm.df_fit_results.columns
    assert list(vm.df_fit_results["Zone"]) == ["zoneA", "zoneB", "zoneC"]


def test_compute_column_from_expression(test_df, qapp, monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    monkeypatch.setattr(QMessageBox, "information", lambda *args, **kwargs: None)
    monkeypatch.setattr(QMessageBox, "warning", lambda *args, **kwargs: None)

    settings = MSettings()
    vm = VMWorkspaceSpectra(settings)
    vm.df_fit_results = test_df.copy()

    vm.compute_column_from_expression("Ratio", "P1_ampli / P2_ampli")
    assert "Ratio" in vm.df_fit_results.columns
    assert list(vm.df_fit_results["Ratio"]) == [2.0, 2.0, 2.0]
